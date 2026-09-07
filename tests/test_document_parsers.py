"""文档解析策略的基础行为测试。"""

from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter

from zq_rag_app.utils.document_parser import parse_document
from zq_rag_app.utils.parsers import (
    DocxDocumentParser,
    DocumentParserFactory,
    ExcelDocumentParser,
    MarkdownDocumentParser,
    PdfDocumentParser,
    TxtDocumentParser,
    UnsupportedDocumentTypeError,
)


def test_factory_selects_strategy_by_extension() -> None:
    factory = DocumentParserFactory()

    assert isinstance(factory.get_parser("manual.pdf"), PdfDocumentParser)
    assert isinstance(factory.get_parser("manual.docx"), DocxDocumentParser)
    assert isinstance(factory.get_parser("manual.txt"), TxtDocumentParser)
    assert isinstance(factory.get_parser("manual.md"), MarkdownDocumentParser)
    assert isinstance(factory.get_parser("manual.xlsx"), ExcelDocumentParser)


def test_factory_rejects_unknown_extension() -> None:
    with pytest.raises(UnsupportedDocumentTypeError):
        DocumentParserFactory().get_parser("manual.pptx")


def test_factory_can_be_created_without_default_strategies() -> None:
    factory = DocumentParserFactory([])

    assert factory.supported_extensions == ()


def test_txt_parser_reads_chinese_text(tmp_path: Path) -> None:
    path = tmp_path / "说明.txt"
    path.write_text("第一段\n\n\n第二段", encoding="utf-8")

    blocks = parse_document(path)

    assert len(blocks) == 1
    assert blocks[0].content == "第一段\n\n第二段"
    assert blocks[0].metadata["encoding"] == "utf-8-sig"


def test_txt_parser_falls_back_to_gb18030(tmp_path: Path) -> None:
    """验证常见的 Windows 中文编码也能被正确解析。"""

    path = tmp_path / "gbk说明.txt"
    path.write_bytes("中文文本".encode("gb18030"))

    blocks = parse_document(path)

    assert blocks[0].content == "中文文本"
    assert blocks[0].metadata["encoding"] == "gb18030"


def test_markdown_parser_splits_sections(tmp_path: Path) -> None:
    path = tmp_path / "guide.md"
    path.write_text("前言\n\n# 安装\n执行命令。\n\n## 配置\n修改配置。", encoding="utf-8")

    blocks = parse_document(path)

    assert [block.section_title for block in blocks] == [None, "安装", "配置"]
    assert blocks[1].content == "执行命令。"
    assert blocks[1].heading_level == 1
    assert blocks[2].heading_level == 2
    assert blocks[2].heading_path == ["安装", "配置"]


def test_markdown_parser_recognizes_setext_heading(tmp_path: Path) -> None:
    """验证 Markdown 的双行 Setext 标题不会被误当成普通正文。"""

    path = tmp_path / "setext.md"
    path.write_text("安装说明\n========\n执行安装命令。", encoding="utf-8")

    blocks = parse_document(path)

    assert len(blocks) == 1
    assert blocks[0].section_title == "安装说明"
    assert blocks[0].content == "执行安装命令。"


def test_docx_parser_keeps_heading_and_table(tmp_path: Path) -> None:
    path = tmp_path / "manual.docx"
    document = Document()
    document.add_heading("产品说明", level=1)
    document.add_paragraph("这是正文。")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "名称"
    table.cell(0, 1).text = "知识库"
    document.save(path)

    blocks = parse_document(path)

    assert len(blocks) == 1
    assert blocks[0].section_title == "产品说明"
    assert blocks[0].heading_level == 1
    assert blocks[0].heading_path == ["产品说明"]
    assert "名称 | 知识库" in blocks[0].content


def test_docx_parser_preserves_nested_heading_path(tmp_path: Path) -> None:
    path = tmp_path / "nested.docx"
    document = Document()
    document.add_heading("第一章", level=1)
    document.add_paragraph("章正文。")
    document.add_heading("第一节", level=2)
    document.add_paragraph("节正文。")
    document.save(path)

    blocks = parse_document(path)

    assert blocks[0].heading_path == ["第一章"]
    assert blocks[1].heading_path == ["第一章", "第一节"]


def test_excel_parser_creates_one_block_per_non_empty_sheet(tmp_path: Path) -> None:
    path = tmp_path / "data.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "产品"
    sheet.append(["名称", "价格"])
    sheet.append(["标准版", 99])
    workbook.create_sheet("空表")
    workbook.save(path)
    workbook.close()

    blocks = parse_document(path)

    assert len(blocks) == 1
    assert blocks[0].section_title == "产品"
    assert blocks[0].content == "名称\t价格\n标准版\t99"


def test_pdf_parser_skips_blank_page(tmp_path: Path) -> None:
    path = tmp_path / "blank.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    with path.open("wb") as output:
        writer.write(output)

    assert parse_document(path) == []


def test_pdf_parser_keeps_page_number_and_ignores_empty_pages(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """隔离 pypdf，验证解析策略对正文、空页和页码的处理。"""

    path = tmp_path / "manual.pdf"
    path.write_bytes(b"%PDF-1.4 unit-test")

    class FakePage:
        def __init__(self, text: str | None) -> None:
            self._text = text

        def extract_text(self) -> str | None:
            return self._text

    class FakeReader:
        def __init__(self, _: str) -> None:
            self.pages = [FakePage("第一页内容"), FakePage(None), FakePage("第三页内容")]

    monkeypatch.setattr(
        "zq_rag_app.utils.parsers.pdf_parser.PdfReader",
        FakeReader,
    )

    blocks = parse_document(path)

    assert [block.content for block in blocks] == ["第一页内容", "第三页内容"]
    assert [block.page_num for block in blocks] == [1, 3]
    assert blocks[1].metadata["page_num"] == 3
