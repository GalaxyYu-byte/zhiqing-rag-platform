"""预览原文件、工作表内容、权限与完整性校验。"""

from hashlib import sha256
from io import BytesIO
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from zhiqing_rag.core.config import Settings
from zhiqing_rag.main import create_app
from zhiqing_rag.modules.documents.preview import preview_content
from zhiqing_rag.modules.documents.router import get_upload_identity, get_upload_service
from zhiqing_rag.modules.documents.schemas import UploadError, UploadIdentity
from zhiqing_rag.modules.documents.service import DocumentUploadService


def workbook_bytes():
    book = Workbook()
    book.active.title = "入职事项"
    book.active.append(["事项", "负责人"])
    book.active.append(["资料登记", "人事"])
    book.create_sheet("复核").append(["<script>alert(1)</script>", "=1+1"])
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


@pytest.fixture
def preview_api():
    content = workbook_bytes()
    identity = UploadIdentity(1, 2, 3, None, 2)
    state = SimpleNamespace(denied=False, content=content, reads=0, released=False)
    row = {
        "original_filename": "入职事项.xlsx",
        "mime_type": "application/octet-stream",
        "bucket": "documents",
        "object_key": "original",
        "file_size": len(content),
        "file_sha256": sha256(content).hexdigest(),
    }

    def revision(current, doc_id, rev_id):
        assert current == identity
        if state.denied or doc_id != 83 or rev_id != 84:
            raise UploadError(404, "DOCUMENT_NOT_FOUND", "文档不存在或没有访问权限")
        return row

    def get_object(bucket, key):
        assert (bucket, key) == ("documents", "original")
        state.reads += 1
        return SimpleNamespace(
            read=lambda size: state.content[:size],
            close=lambda: None,
            release_conn=lambda: setattr(state, "released", True),
        )

    service = DocumentUploadService(
        SimpleNamespace(preview_revision=revision),
        SimpleNamespace(client=SimpleNamespace(get_object=get_object)),
        Settings(_env_file=None, app_env="test"),
    )
    app = create_app()
    app.dependency_overrides[get_upload_service] = lambda: service
    app.dependency_overrides[get_upload_identity] = lambda: identity
    with TestClient(app) as client:
        yield client, state, content


def test_excel_preview_preserves_sheets_tables_and_formula_warnings(preview_api):
    client, state, _ = preview_api
    response = client.get("/api/documents/83/preview?revision_id=84")
    assert response.status_code == 200
    body = response.json()
    assert body["elements"][0]["sheet_name"] == "入职事项"
    assert body["elements"][0]["table_rows"] == [["事项", "负责人"], ["资料登记", "人事"]]
    assert body["elements"][1]["sheet_name"] == "复核"
    assert body["elements"][1]["table_rows"][0] == ["<script>alert(1)</script>", "=1+1"]
    assert body["warnings"] and not body["truncated"]
    assert response.headers["cache-control"] == "no-store"
    assert state.released


def test_original_download_returns_exact_bytes(preview_api):
    client, _, content = preview_api
    response = client.get("/api/documents/83/content?revision_id=84")
    assert response.status_code == 200 and response.content == content
    assert response.headers["content-disposition"].startswith("attachment; filename*=UTF-8''")
    assert response.headers["x-content-type-options"] == "nosniff"


@pytest.mark.parametrize("endpoint", ["preview", "content"])
def test_preview_rejects_unauthorized_or_mismatched_revisions_before_storage(preview_api, endpoint):
    client, state, _ = preview_api
    assert client.get(f"/api/documents/83/{endpoint}?revision_id=999").status_code == 404
    state.denied = True
    assert client.get(f"/api/documents/83/{endpoint}?revision_id=84").status_code == 404
    assert state.reads == 0


def test_preview_checks_original_integrity(preview_api):
    client, state, _ = preview_api
    state.content = b"corrupted"
    response = client.get("/api/documents/83/preview?revision_id=84")
    assert response.status_code == 409
    assert response.json()["detail"]["code"] == "FILE_INTEGRITY_FAILED"


def test_pending_file_preview_does_not_read_or_persist_objects(preview_api):
    client, state, content = preview_api
    response = client.post("/api/documents/preview", files={"file": ("待上传.xlsx", content)})
    assert response.status_code == 200 and response.json()["elements"][0]["table_rows"]
    assert state.reads == 0


def test_large_preview_reports_truncation():
    book = Workbook()
    for index in range(205):
        book.active.append([str(index)])
    stream = BytesIO()
    book.save(stream)
    body = preview_content(stream.getvalue(), "large.xlsx", 1024 * 1024)
    assert body["truncated"] and len(body["elements"][0]["table_rows"]) == 200
