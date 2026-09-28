"""离线文档格式校验的真实文件与伪装文件测试。"""

import hashlib
import zipfile
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter

from zhiqing_rag.document_processing.format_validation import (
    DocumentFormat,
    DocumentFormatError,
    validate_document_format,
)


def test_validate_supported_documents(tmp_path: Path) -> None:
    pdf_path = tmp_path / "report.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=400)
    with pdf_path.open("wb") as file:
        writer.write(file)

    docx_path = tmp_path / "manual.docx"
    document = Document()
    document.add_paragraph("测试文档")
    document.save(docx_path)

    xlsx_path = tmp_path / "data.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "测试数据"
    workbook.save(xlsx_path)

    markdown_path = tmp_path / "guide.md"
    markdown_path.write_text("# 标题\n正文", encoding="utf-8")
    txt_path = tmp_path / "notes.txt"
    txt_path.write_bytes("说明文字".encode("gb18030"))

    for path, expected in (
        (pdf_path, DocumentFormat.PDF),
        (docx_path, DocumentFormat.DOCX),
        (xlsx_path, DocumentFormat.XLSX),
        (markdown_path, DocumentFormat.MARKDOWN),
        (txt_path, DocumentFormat.TXT),
    ):
        result = validate_document_format(path)
        assert result.format == expected
        assert result.file_size == path.stat().st_size
        assert result.file_sha256 == hashlib.sha256(path.read_bytes()).hexdigest()

    assert validate_document_format(txt_path).text_encoding == "gb18030"


def test_validation_uses_original_filename_for_temporary_file(tmp_path: Path) -> None:
    path = tmp_path / "upload.tmp"
    path.write_text("正文", encoding="utf-8")

    result = validate_document_format(path, original_filename="document.md")

    assert result.format == DocumentFormat.MARKDOWN


@pytest.mark.parametrize(
    ("filename", "content", "expected_code"),
    [
        ("fake.pdf", b"not a pdf", "FILE_TYPE_MISMATCH"),
        ("broken.pdf", b"%PDF-1.7\nnot a complete pdf", "INVALID_FILE"),
        ("fake.txt", b"%PDF-1.7\n", "FILE_TYPE_MISMATCH"),
        ("fake.md", b"abc\x00def", "INVALID_FILE"),
        ("broken.docx", b"PK\x03\x04not a complete zip", "INVALID_FILE"),
        ("fake.xls", b"not an xls", "FILE_TYPE_MISMATCH"),
        ("legacy.doc", b"content", "UNSUPPORTED_FORMAT"),
        ("empty.md", b"", "FILE_EMPTY"),
    ],
)
def test_rejects_invalid_files(
    tmp_path: Path, filename: str, content: bytes, expected_code: str
) -> None:
    path = tmp_path / filename
    path.write_bytes(content)

    with pytest.raises(DocumentFormatError) as error:
        validate_document_format(path)

    assert error.value.code == expected_code


def test_rejects_ooxml_with_wrong_document_type(tmp_path: Path) -> None:
    path = tmp_path / "fake.docx"
    workbook = Workbook()
    workbook.save(path)

    with pytest.raises(DocumentFormatError, match="类型与扩展名不一致") as error:
        validate_document_format(path)

    assert error.value.code == "FILE_TYPE_MISMATCH"


def test_rejects_archive_expansion_bomb(tmp_path: Path) -> None:
    path = tmp_path / "bomb.docx"
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", "x" * 100_000)

    with pytest.raises(DocumentFormatError) as error:
        validate_document_format(path)

    assert error.value.code == "INVALID_FILE"


def test_rejects_size_and_mime_mismatch(tmp_path: Path) -> None:
    path = tmp_path / "document.txt"
    path.write_text("正文", encoding="utf-8")

    with pytest.raises(DocumentFormatError) as size_error:
        validate_document_format(path, max_file_size=1)
    assert size_error.value.code == "FILE_TOO_LARGE"

    with pytest.raises(DocumentFormatError) as mime_error:
        validate_document_format(path, declared_mime_type="application/pdf")
    assert mime_error.value.code == "FILE_TYPE_MISMATCH"
