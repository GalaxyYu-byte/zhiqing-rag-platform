"""Excel 工作簿解析策略。"""

from datetime import date, datetime, time
from pathlib import Path
from typing import Any

import xlrd
from openpyxl import load_workbook

from .base import DocumentParseError, DocumentParser, ParsedBlock
from .common import base_metadata, clean_text


def _value_to_text(value: Any) -> str:
    """将 openpyxl 单元格值稳定转换为文本。"""

    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value)


def _trim_trailing_empty_cells(values: list[str]) -> list[str]:
    """删除行尾无意义空单元格，同时保留中间空列的表格位置。"""

    while values and not values[-1]:
        values.pop()
    return values


class ExcelDocumentParser(DocumentParser):
    """解析新版和旧版 Excel 文件。

    ``.xlsx/.xlsm`` 使用 openpyxl，旧版二进制 ``.xls`` 使用 xlrd。每个非空
    Sheet 生成一个 ``ParsedBlock``，单元格用制表符连接，便于后续保留行列
    关系。``data_only=True`` 会优先读取公式的缓存计算值，而不是公式字符串。
    """

    supported_extensions = frozenset({".xlsx", ".xlsm", ".xltx", ".xltm", ".xls"})

    def parse(self, file_path: str | Path) -> list[ParsedBlock]:
        path = self.validate_file(file_path)
        if path.suffix.lower() == ".xls":
            return self._parse_xls(path)
        return self._parse_openxml(path)

    def _parse_openxml(self, path: Path) -> list[ParsedBlock]:
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:
            raise DocumentParseError(f"Excel 文件打开失败: {path}") from exc

        blocks: list[ParsedBlock] = []
        try:
            for sheet_index, sheet in enumerate(workbook.worksheets, start=1):
                lines: list[str] = []
                for row in sheet.iter_rows(values_only=True):
                    values = _trim_trailing_empty_cells(
                        [_value_to_text(value) for value in row]
                    )
                    if any(values):
                        lines.append("\t".join(values))

                self._append_sheet_block(
                    blocks, path, sheet.title, sheet_index, lines
                )
        finally:
            # read_only 工作簿可能持有压缩包句柄，必须显式关闭。
            workbook.close()

        return blocks

    def _parse_xls(self, path: Path) -> list[ParsedBlock]:
        try:
            workbook = xlrd.open_workbook(str(path), on_demand=True)
        except Exception as exc:
            raise DocumentParseError(f"XLS 文件打开失败: {path}") from exc

        blocks: list[ParsedBlock] = []
        try:
            for sheet_index, sheet in enumerate(workbook.sheets(), start=1):
                lines: list[str] = []
                for row_index in range(sheet.nrows):
                    values = [
                        self._xls_cell_to_text(
                            sheet.cell(row_index, column), workbook.datemode
                        )
                        for column in range(sheet.ncols)
                    ]
                    values = _trim_trailing_empty_cells(values)
                    if any(values):
                        lines.append("\t".join(values))

                self._append_sheet_block(
                    blocks, path, sheet.name, sheet_index, lines
                )
        finally:
            workbook.release_resources()

        return blocks

    @staticmethod
    def _xls_cell_to_text(cell: xlrd.sheet.Cell, datemode: int) -> str:
        """根据 xlrd 单元格类型转换旧版 Excel 的值。"""

        if cell.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK}:
            return ""
        if cell.ctype == xlrd.XL_CELL_DATE:
            return xlrd.xldate_as_datetime(cell.value, datemode).isoformat()
        if cell.ctype == xlrd.XL_CELL_BOOLEAN:
            return "TRUE" if cell.value else "FALSE"
        if cell.ctype == xlrd.XL_CELL_ERROR:
            return f"#ERROR({int(cell.value)})"
        return str(cell.value)

    @staticmethod
    def _append_sheet_block(
        blocks: list[ParsedBlock],
        path: Path,
        sheet_name: str,
        sheet_index: int,
        lines: list[str],
    ) -> None:
        content = clean_text("\n".join(lines))
        if not content:
            return

        blocks.append(
            ParsedBlock(
                content=content,
                section_title=sheet_name,
                metadata={
                    **base_metadata(path, ExcelDocumentParser.__name__),
                    "sheet_name": sheet_name,
                    "sheet_index": sheet_index,
                },
            )
        )

