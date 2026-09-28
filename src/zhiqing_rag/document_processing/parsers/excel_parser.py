"""保留 Sheet、单元格范围和表格行列的 Excel 解析策略。"""

from contextlib import contextmanager
from datetime import date, datetime, time
from itertools import zip_longest
from pathlib import Path
from typing import Any

import xlrd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from zhiqing_rag.document_processing.format_validation import (
    DocumentFormat,
    ValidatedDocument,
)

from .base import DocumentParser
from .models import ElementKind, ParsedElement, ParseResult, ParseWarning

MAX_TABLE_ROWS = 500


def _value_to_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)


class _TableRegion:
    def __init__(self, sheet_name: str, elements: list[ParsedElement]) -> None:
        self.sheet_name = sheet_name
        self.elements = elements
        self.rows: list[tuple[int, dict[int, str]]] = []
        self.min_column: int | None = None
        self.max_column: int | None = None

    def add(self, row_number: int, cells: dict[int, str]) -> None:
        if not cells:
            self.flush()
            return
        self.rows.append((row_number, cells))
        first = min(cells)
        last = max(cells)
        self.min_column = first if self.min_column is None else min(self.min_column, first)
        self.max_column = last if self.max_column is None else max(self.max_column, last)
        if len(self.rows) >= MAX_TABLE_ROWS:
            self.flush()

    def flush(self) -> None:
        if not self.rows or self.min_column is None or self.max_column is None:
            return
        rows = tuple(
            tuple(cells.get(column, "") for column in range(self.min_column, self.max_column + 1))
            for _, cells in self.rows
        )
        start_row = self.rows[0][0]
        end_row = self.rows[-1][0]
        cell_range = (
            f"{get_column_letter(self.min_column)}{start_row}:"
            f"{get_column_letter(self.max_column)}{end_row}"
        )
        self.elements.append(
            ParsedElement(
                kind=ElementKind.TABLE,
                text="\n".join("\t".join(row) for row in rows),
                order=len(self.elements),
                heading_path=(self.sheet_name,),
                sheet_name=self.sheet_name,
                cell_range=cell_range,
                table_rows=rows,
                source_ref=f"sheet:{self.sheet_name}!{cell_range}",
            )
        )
        self.rows.clear()
        self.min_column = None
        self.max_column = None


@contextmanager
def _xlsx_workbook(path: Path, *, data_only: bool):
    # 后台下载使用 .tmp 路径；格式已校验，通过文件流读取避免扩展名限制。
    with path.open("rb") as source:
        workbook = load_workbook(source, read_only=True, data_only=data_only)
        try:
            yield workbook
        finally:
            workbook.close()


class XlsxParser(DocumentParser):
    formats = frozenset({DocumentFormat.XLSX})

    def parse(self, path: Path, validated: ValidatedDocument) -> ParseResult:
        del validated
        elements: list[ParsedElement] = []
        warnings: list[ParseWarning] = []
        with (
            _xlsx_workbook(path, data_only=True) as values_book,
            _xlsx_workbook(path, data_only=False) as formulas_book,
        ):
            for values_sheet in values_book.worksheets:
                formulas_sheet = formulas_book[values_sheet.title]
                region = _TableRegion(values_sheet.title, elements)
                value_rows = values_sheet.iter_rows()
                formula_rows = formulas_sheet.iter_rows()
                for row_number, (value_row, formula_row) in enumerate(
                    zip_longest(value_rows, formula_rows, fillvalue=()), start=1
                ):
                    cells: dict[int, str] = {}
                    for column in range(1, max(len(value_row), len(formula_row)) + 1):
                        value_cell = value_row[column - 1] if column <= len(value_row) else None
                        formula_cell = (
                            formula_row[column - 1] if column <= len(formula_row) else None
                        )
                        value = getattr(value_cell, "value", None)
                        formula = getattr(formula_cell, "value", None)
                        if getattr(formula_cell, "data_type", None) == "f" and value is None:
                            coordinate = f"{get_column_letter(column)}{row_number}"
                            warnings.append(
                                ParseWarning(
                                    code="FORMULA_CACHE_MISSING",
                                    message="公式缺少已计算的缓存值，已保留公式表达式",
                                    sheet_name=values_sheet.title,
                                    cell_range=coordinate,
                                )
                            )
                            value = formula
                        text = _value_to_text(value)
                        if text:
                            cells[column] = text
                    region.add(row_number, cells)
                region.flush()
        return ParseResult(tuple(elements), tuple(warnings))


def _xls_cell_to_text(cell: xlrd.sheet.Cell, datemode: int) -> str:
    if cell.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK}:
        return ""
    if cell.ctype == xlrd.XL_CELL_DATE:
        return xlrd.xldate_as_datetime(cell.value, datemode).isoformat()
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return "TRUE" if cell.value else "FALSE"
    if cell.ctype == xlrd.XL_CELL_ERROR:
        return f"#ERROR({int(cell.value)})"
    return _value_to_text(cell.value)


class XlsParser(DocumentParser):
    formats = frozenset({DocumentFormat.XLS})

    def parse(self, path: Path, validated: ValidatedDocument) -> ParseResult:
        del validated
        workbook = xlrd.open_workbook(str(path), on_demand=True)
        elements: list[ParsedElement] = []
        try:
            for sheet in workbook.sheets():
                region = _TableRegion(sheet.name, elements)
                for row_index in range(sheet.nrows):
                    cells = {
                        column + 1: text
                        for column in range(sheet.ncols)
                        if (
                            text := _xls_cell_to_text(
                                sheet.cell(row_index, column), workbook.datemode
                            )
                        )
                    }
                    region.add(row_index + 1, cells)
                region.flush()
        finally:
            workbook.release_resources()

        return ParseResult(tuple(elements))
